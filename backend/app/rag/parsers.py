
"""
parsers.py
-----------
Extract readable text from supported file types.

Each parser returns Documents with source, page, and section metadata.
parse_file() also returns notes about skipped or removed content.

Validation includes extension allow-listing, magic-byte checks, ZIP size
limits, unsafe internal ZIP paths, and an extracted-text size limit.
"""

import io
import os
import re
import zipfile

from langchain_core.documents import Document

from app.config import (
    MAX_EXTRACTED_CHARS,
    MAX_XLSX_ROWS_PER_SHEET,
    MAX_ZIP_UNCOMPRESSED_MB,
    logger,
)

SUPPORTED_EXTENSIONS = (
    ".pdf", ".txt", ".md", ".docx", ".html", ".htm", ".xlsx"
)
_XLSX_ROWS_PER_BLOCK = 25


# --------------------------------------------------------------------------
# Validation
# --------------------------------------------------------------------------

def _check_magic(ext: str, data: bytes) -> None:
    if ext == ".pdf" and not data.lstrip()[:5].startswith(b"%PDF"):
        raise ValueError("File is named .pdf but is not a PDF")

    if ext in (".docx", ".xlsx") and not data.startswith(b"PK"):
        raise ValueError(f"File is named {ext} but is not a valid Office file")

    if ext in (".txt", ".md", ".html", ".htm") and b"\x00" in data[:4096]:
        raise ValueError(f"File is named {ext} but looks like binary data")


def _check_zip(data: bytes) -> None:
    """Validate Office ZIP contents without leaving the archive open."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()

            total_uncompressed = sum(
                entry.file_size for entry in entries
            )
            max_uncompressed = MAX_ZIP_UNCOMPRESSED_MB * 1024 * 1024

            if total_uncompressed > max_uncompressed:
                raise ValueError(
                    "File expands to an unreasonable size (possible zip bomb)"
                )

            for entry in entries:
                path_parts = entry.filename.replace("\\", "/").split("/")
                if entry.filename.startswith("/") or ".." in path_parts:
                    raise ValueError("File contains unsafe internal paths")

    except zipfile.BadZipFile as exc:
        raise ValueError("Corrupt or invalid Office file") from exc


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16"):
        try:
            return data.decode(encoding)
        except UnicodeError:
            continue

    return data.decode("latin-1")


def _doc(
    text: str,
    source: str,
    page: int,
    section: str = "",
    **extra,
) -> Document:
    return Document(
        page_content=text,
        metadata={
            "source": source,
            "page": page,
            "section": section,
            **extra,
        },
    )


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------

def _parse_pdf(
    name: str,
    data: bytes,
) -> tuple[list[Document], list[str]]:
    import fitz
    from pypdf import PdfReader

    notes: list[str] = []

    try:
        pdf = fitz.open(stream=data, filetype="pdf")
        try:
            if pdf.needs_pass:
                raise ValueError("PDF is password-protected")

            pages = [
                (page_number, page.get_text() or "")
                for page_number, page in enumerate(pdf, start=1)
            ]
        finally:
            pdf.close()

    except ValueError:
        raise
    except Exception:
        try:
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise ValueError("PDF is password-protected")

            pages = [
                (page_number, page.extract_text() or "")
                for page_number, page in enumerate(reader.pages, start=1)
            ]
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("Could not read this PDF (corrupt?)") from exc

    docs = [
        _doc(text, name, page_number)
        for page_number, text in pages
        if text.strip()
    ]

    empty_pages = len(pages) - len(docs)
    if empty_pages:
        notes.append(
            f"{empty_pages} of {len(pages)} pages had no extractable text "
            "(scanned? OCR is not supported yet)"
        )

    return docs, notes


# --------------------------------------------------------------------------
# TXT / Markdown
# --------------------------------------------------------------------------

def _parse_txt(
    name: str,
    data: bytes,
) -> tuple[list[Document], list[str]]:
    text = _decode(data)
    docs = [_doc(text, name, 1)] if text.strip() else []
    return docs, []


def _parse_md(
    name: str,
    data: bytes,
) -> tuple[list[Document], list[str]]:
    """Split Markdown at headings while ignoring heading-like text in fences."""
    sections: list[tuple[str, list[str]]] = [("", [])]
    in_fence = False

    for line in _decode(data).splitlines():
        if line.strip().startswith("```"):
            in_fence = not in_fence

        heading_match = (
            None
            if in_fence
            else re.match(r"^\s{0,3}(#{1,6})\s+(.*?)\s*#*\s*$", line)
        )

        if heading_match:
            sections.append((heading_match.group(2), [line]))
        else:
            sections[-1][1].append(line)

    docs = []
    for page_number, (heading, lines) in enumerate(sections, start=1):
        body = "\n".join(lines).strip()
        if body:
            docs.append(_doc(body, name, page_number, heading))

    return docs, []


# --------------------------------------------------------------------------
# DOCX
# --------------------------------------------------------------------------

def _parse_docx(
    name: str,
    data: bytes,
) -> tuple[list[Document], list[str]]:
    _check_zip(data)

    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ValueError(
            "Could not read this Word file (corrupt or password-protected?)"
        ) from exc

    sections: list[tuple[str, list[str]]] = [("", [])]

    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]

        if tag == "p":
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if not text:
                continue

            style = (
                paragraph.style.name or ""
                if paragraph.style is not None
                else ""
            )

            if (
                style.lower().startswith("heading")
                or style.lower() == "title"
            ):
                sections.append((text, [text]))
            else:
                sections[-1][1].append(text)

        elif tag == "tbl":
            table = Table(child, document)
            for row in table.rows:
                cells = []
                for cell in row.cells:
                    cell_text = cell.text.strip().replace("\n", " ")
                    if cell_text and (
                        not cells or cells[-1] != cell_text
                    ):
                        cells.append(cell_text)

                if cells:
                    sections[-1][1].append(" | ".join(cells))

    docs = []
    for page_number, (heading, lines) in enumerate(sections, start=1):
        body = "\n".join(lines).strip()
        if body:
            docs.append(_doc(body, name, page_number, heading))

    return docs, []


# --------------------------------------------------------------------------
# HTML
# --------------------------------------------------------------------------

_HIDDEN_STYLE = re.compile(
    r"display\s*:\s*none|visibility\s*:\s*hidden|font-size\s*:\s*0|"
    r"opacity\s*:\s*0(\.0+)?\b|"
    r"(left|top|text-indent)\s*:\s*-\d{3,}",
    re.I,
)


def _parse_html(
    name: str,
    data: bytes,
) -> tuple[list[Document], list[str]]:
    from bs4 import BeautifulSoup, Comment

    soup = BeautifulSoup(_decode(data), "html.parser")
    notes: list[str] = []

    for tag in soup(
        [
            "script", "style", "noscript", "template",
            "iframe", "object", "embed", "svg",
        ]
    ):
        tag.decompose()

    for comment in soup.find_all(
        string=lambda value: isinstance(value, Comment)
    ):
        comment.extract()

    hidden_elements = [
        tag
        for tag in soup.find_all(True)
        if (
            tag.has_attr("hidden")
            or tag.get("aria-hidden") == "true"
            or _HIDDEN_STYLE.search(tag.get("style", "") or "")
        )
    ]

    removed = 0
    for tag in hidden_elements:
        if tag.parent is None:
            continue

        if tag.get_text(strip=True):
            removed += 1
        tag.decompose()

    if removed:
        notes.append(
            f"removed {removed} hidden HTML element(s) containing text"
        )

    sections: list[tuple[str, list[str]]] = [("", [])]
    root = soup.body or soup

    for element in root.find_all(
        [
            "h1", "h2", "h3", "p", "li", "tr",
            "pre", "blockquote", "td", "th",
        ]
    ):
        if element.name in ("td", "th") and element.find_parent("tr"):
            continue

        if element.name in ("p", "li", "blockquote") and element.find_parent(
            ["li", "blockquote", "tr"]
        ):
            continue

        if element.name == "tr":
            cells = [
                cell.get_text(" ", strip=True)
                for cell in element.find_all(["td", "th"])
            ]
            text = " | ".join(cell for cell in cells if cell)
        else:
            text = element.get_text(" ", strip=True)

        if not text:
            continue

        if element.name in ("h1", "h2", "h3"):
            sections.append((text, [text]))
        else:
            sections[-1][1].append(text)

    docs = []
    for page_number, (heading, lines) in enumerate(sections, start=1):
        body = "\n".join(lines).strip()
        if body:
            docs.append(_doc(body, name, page_number, heading))

    if not docs:
        fallback_text = root.get_text("\n", strip=True)
        if fallback_text:
            docs.append(_doc(fallback_text, name, 1))

    return docs, notes


# --------------------------------------------------------------------------
# XLSX
# --------------------------------------------------------------------------

def _parse_xlsx(
    name: str,
    data: bytes,
) -> tuple[list[Document], list[str]]:
    _check_zip(data)

    import openpyxl

    try:
        workbook = openpyxl.load_workbook(
            io.BytesIO(data),
            read_only=True,
            data_only=True,
        )
    except Exception as exc:
        raise ValueError(
            "Could not read this Excel file (corrupt or password-protected?)"
        ) from exc

    notes: list[str] = []
    docs: list[Document] = []
    document_number = 0

    try:
        for worksheet in workbook.worksheets:
            if getattr(worksheet, "sheet_state", "visible") != "visible":
                notes.append(f"skipped hidden sheet '{worksheet.title}'")
                continue

            header: list[str] = []
            block: list[str] = []
            block_start = 1
            rows_seen = 0
            last_row = 1

            def flush(end_row: int) -> None:
                nonlocal block, block_start, document_number

                if block:
                    document_number += 1
                    docs.append(
                        _doc(
                            "\n".join(block),
                            name,
                            document_number,
                            f"Sheet: {worksheet.title} "
                            f"(rows {block_start}-{end_row})",
                        )
                    )
                block = []

            for row_index, row in enumerate(
                worksheet.iter_rows(values_only=True),
                start=1,
            ):
                if row_index > MAX_XLSX_ROWS_PER_SHEET:
                    notes.append(
                        f"sheet '{worksheet.title}' truncated at "
                        f"{MAX_XLSX_ROWS_PER_SHEET} rows"
                    )
                    break

                last_row = row_index
                cells = [
                    "" if value is None else str(value).strip()
                    for value in row
                ]

                if not any(cells):
                    continue

                rows_seen += 1

                if not header:
                    header = [
                        cell or f"col{index + 1}"
                        for index, cell in enumerate(cells)
                    ]
                    block_start = row_index
                    block.append("Columns: " + " | ".join(header))
                    continue

                line = "; ".join(
                    f"{column}: {cell}"
                    for column, cell in zip(header, cells)
                    if cell
                )

                if line:
                    block.append(line)

                if len(block) >= _XLSX_ROWS_PER_BLOCK:
                    flush(row_index)
                    block_start = row_index + 1

            # last_row is initialized, so empty sheets are safe.
            flush(last_row if rows_seen else 1)

    finally:
        workbook.close()

    return docs, notes


_PARSERS = {
    ".pdf": _parse_pdf,
    ".txt": _parse_txt,
    ".md": _parse_md,
    ".docx": _parse_docx,
    ".html": _parse_html,
    ".htm": _parse_html,
    ".xlsx": _parse_xlsx,
}


def parse_file(
    file_name: str,
    data: bytes,
) -> tuple[list[Document], list[str]]:
    extension = os.path.splitext(file_name)[1].lower()

    if extension not in _PARSERS:
        logger.warning("Rejected unsupported file type: %s", extension)
        supported = ", ".join(SUPPORTED_EXTENSIONS)
        raise ValueError(
            f"Unsupported file type: {extension or '(none)'}. "
            f"Supported: {supported}"
        )

    if not data:
        raise ValueError("File is empty")

    _check_magic(extension, data)
    documents, notes = _PARSERS[extension](file_name, data)

    if not documents:
        raise ValueError("No extractable text found in this file")

    total_chars = sum(len(document.page_content) for document in documents)

    if total_chars > MAX_EXTRACTED_CHARS:
        raise ValueError(
            f"Document has too much text ({total_chars:,} characters; "
            f"limit {MAX_EXTRACTED_CHARS:,})"
        )

    logger.info(
        "Parsed %s: %d section(s), %d chars",
        file_name,
        len(documents),
        total_chars,
    )
    return documents, notes
